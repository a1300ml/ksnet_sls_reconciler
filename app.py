from __future__ import annotations

import json
import re
import shutil
import sys
import tkinter as tk
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


APP_TITLE = "런케이에듀 SLS 매출 조회"
FILE_TYPES = [("SLS Excel", "*.xlsx *.xlsm"), ("모든 파일", "*.*")]
APP_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
PROJECTS_DIR = APP_DIR / "프로젝트"
SOURCES = ("현금매출", "신용카드매출", "판매(결제)대행매출")
REFERENCE_SOURCES = ("현금매출", "판매(결제)대행매출", "신용카드매출")
CATEGORIES = ("수강료", "가맹점유저", "교재매출", "가맹점매출")
ACADEMY_ORDER = (
    "광명철산 직영관", "광진 직영관", "구리 직영관", "국풍온", "런케이에듀 CP",
    "목동 직영관", "별내 직영관", "송파 직영관", "의정부민락 직영관",
    "일산백마 직영관", "중계 직영관", "H&C 평생교육원 노원",
)


@dataclass(frozen=True)
class SalesRow:
    year: int
    month: int
    academy: str
    category: str
    source: str
    amount: int
    sheet: str
    row: int


@dataclass(frozen=True)
class ReferenceRow:
    year: int
    month: int
    source: str
    amount: int
    file: str


@dataclass(frozen=True)
class AllocatedRow:
    year: int
    month: int
    source: str
    academy: str
    category: str
    amount: int


def amount(value) -> int:
    if value in (None, ""):
        return 0
    try:
        return int(round(float(str(value).replace(",", ""))))
    except (TypeError, ValueError):
        return 0


def normalize(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def period_from_sheet(ws) -> tuple[int, int]:
    title = normalize(ws.cell(1, 1).value)
    match = re.search(r"(\d{4})년\s*(\d{1,2})월", title)
    if match:
        return int(match.group(1)), int(match.group(2))
    match = re.search(r"(?:(\d{4})\D*)?(\d{1,2})월", ws.title)
    if match:
        return int(match.group(1) or 0), int(match.group(2))
    raise ValueError(f"'{ws.title}' 시트에서 조회 월을 찾지 못했습니다.")


def runk_category(academy: str, item: str) -> str | None:
    """런케이에듀에 귀속되는 행만 반환하고 나머지는 제외한다."""
    academy = normalize(academy)
    item = normalize(item)
    if academy == "런케이에듀 CP":
        if item == "수강료":
            return "가맹점유저"
        if "교재" in item:
            return "교재매출"
        if item in {"컨텐츠", "콘텐츠", "기타"}:
            return "가맹점매출"
        return None

    # 모든 관의 교재비와 콘텐츠비는 런케이에듀 매출이다.
    if "교재" in item:
        return "교재매출"
    if item in {"컨텐츠", "콘텐츠"}:
        return "가맹점유저"

    # 수강료는 국풍온과 H&C 평생교육원 노원만 런케이에듀에 귀속한다.
    if academy in {"국풍온", "H&C 평생교육원 노원"} and item == "수강료":
        return "수강료"
    return None


def _header_layout(ws) -> tuple[int, dict[str, int]]:
    """결제수단별 '순수납' 열을 찾는다. 병합 셀도 왼쪽 제목을 이어서 읽는다."""
    for top_row in range(1, min(ws.max_row, 12) + 1):
        if normalize(ws.cell(top_row, 1).value) != "학원":
            continue
        sub_row = top_row + 1
        layout: dict[str, int] = {}
        current_group = ""
        for col in range(3, ws.max_column + 1):
            group = normalize(ws.cell(top_row, col).value)
            if group:
                current_group = group
            if normalize(ws.cell(sub_row, col).value) == "순수납":
                if current_group in {"현금/계좌", "창구카드", "가상계좌", "이카드"}:
                    layout[current_group] = col
        if set(layout) == {"현금/계좌", "창구카드", "가상계좌", "이카드"}:
            return sub_row + 1, layout
    raise ValueError(f"'{ws.title}' 시트에서 SLS 집계표 제목을 찾지 못했습니다.")


def read_sls(path: str) -> list[SalesRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows: list[SalesRow] = []
    fallback_year = 0
    for ws in wb.worksheets:
        year, month = period_from_sheet(ws)
        if year:
            fallback_year = year
        elif fallback_year:
            year = fallback_year
        else:
            raise ValueError(f"'{ws.title}' 시트에서 연도를 찾지 못했습니다.")
        first_data_row, cols = _header_layout(ws)
        for row_no in range(first_data_row, ws.max_row + 1):
            academy = normalize(ws.cell(row_no, 1).value)
            item = normalize(ws.cell(row_no, 2).value)
            category = runk_category(academy, item)
            if category is None:
                continue
            values = {
                "현금매출": amount(ws.cell(row_no, cols["현금/계좌"]).value)
                + amount(ws.cell(row_no, cols["가상계좌"]).value),
                "신용카드매출": amount(ws.cell(row_no, cols["창구카드"]).value),
                "판매(결제)대행매출": amount(ws.cell(row_no, cols["이카드"]).value),
            }
            for source, value in values.items():
                rows.append(SalesRow(year, month, academy, category, source, value, ws.title, row_no))
    if not rows:
        raise ValueError("런케이에듀 매출에 해당하는 행이 없습니다.")
    return rows


def _period_from_text(text: str) -> tuple[int, int] | None:
    match = re.search(r"(20\d{2})[./년-]?\s*(\d{1,2})(?:월)?", text)
    if match:
        return int(match.group(1)), int(match.group(2))
    return None


def _parse_reference_pdf(path: str, source: str) -> list[ReferenceRow]:
    text = "\n".join((page.extract_text() or "") for page in PdfReader(path).pages)
    rows = []
    pattern = re.compile(
        r"(?:신용카드|계좌이체|가상계좌)\s+(20\d{2})(\d{2})\s+[\d,]+\s+(-?[\d,]+)"
    )
    for match in pattern.finditer(text):
        rows.append(ReferenceRow(
            int(match.group(1)), int(match.group(2)), source,
            int(match.group(3).replace(",", "")), Path(path).name,
        ))
    if not rows:
        raise ValueError("PDF에서 월별 거래금액 표를 찾지 못했습니다.")
    return rows


def _xlsx_values(path: str) -> list[list]:
    # 일부 KCP 파일은 dimension을 A1로 잘못 기록하므로 일반 모드로 읽는다.
    ws = load_workbook(path, read_only=False, data_only=True).active
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _parse_reference_xlsx(path: str, source: str) -> list[ReferenceRow]:
    values = _xlsx_values(path)
    filename = Path(path).name
    for index, row in enumerate(values):
        headers = [normalize(value) for value in row]

        # KSNET 마이장부: 카드사별 승인금액을 년월별 합산
        if "년월" in headers and "승인금액" in headers:
            period_col, amount_col = headers.index("년월"), headers.index("승인금액")
            grouped = defaultdict(int)
            current_period = None
            credit_totals_found = False
            for data in values[index + 1:]:
                period = _period_from_text(normalize(data[period_col]))
                if period:
                    current_period = period
                # 신용카드 첨부자료에서는 현금영수증이 제외된 월별 신용합계만 사용한다.
                # 카드사 상세와 월별합계를 다시 더하면 중복되므로 신용합계 행만 읽는다.
                if len(data) > 1 and normalize(data[1]) == "신용합계" and current_period:
                    grouped[current_period] = amount(data[amount_col])
                    credit_totals_found = True
            if not credit_totals_found:
                # 다른 형식의 마이장부에는 신용합계 행이 없을 수 있어 카드사 승인금액만 합산한다.
                grouped.clear()
                for data in values[index + 1:]:
                    period = _period_from_text(normalize(data[period_col]))
                    merchant = normalize(data[3]) if len(data) > 3 else ""
                    if period and merchant and merchant != "현금영수증":
                        grouped[period] += amount(data[amount_col])
            return [ReferenceRow(y, m, source, value, filename) for (y, m), value in sorted(grouped.items())]

        # KCP 가상계좌/계좌이체: 월별 거래금액 행
        if "조회년월" in headers and "거래금액" in headers:
            period_col, amount_col = headers.index("조회년월"), headers.index("거래금액")
            rows = []
            for data in values[index + 1:]:
                period = _period_from_text(normalize(data[period_col]))
                if period:
                    rows.append(ReferenceRow(period[0], period[1], source, amount(data[amount_col]), filename))
            return rows

        # KCP 신용카드: 대상기간의 카드사별 매출금액 합계
        card_header = next((h for h in ("매출금액 - 매출대상금액", "신용카드 매출금액") if h in headers), None)
        if "카드사" in headers and card_header:
            amount_col = headers.index(card_header)
            period = next(
                (_period_from_text(normalize(cell)) for prior in values[:index] for cell in prior
                 if _period_from_text(normalize(cell))),
                None,
            )
            if not period:
                period = _period_from_text(filename)
            if not period:
                raise ValueError("KCP 파일에서 대상기간을 찾지 못했습니다.")
            total = 0
            for data in values[index + 1:]:
                if normalize(data[0]) == "합계":
                    break
                total += amount(data[amount_col])
            return [ReferenceRow(period[0], period[1], source, total, filename)]
    raise ValueError("지원하는 부가세 참고자료 형식을 찾지 못했습니다.")


def _parse_reference_xls(path: str, source: str) -> list[ReferenceRow]:
    ws = xlrd.open_workbook(path).sheet_by_index(0)
    values = [ws.row_values(row) for row in range(ws.nrows)]
    filename = Path(path).name
    for index, row in enumerate(values):
        headers = [normalize(value) for value in row]
        if "매출일시" not in headers or "총금액" not in headers:
            continue
        date_col, amount_col = headers.index("매출일시"), headers.index("총금액")
        grouped = defaultdict(int)
        for data in values[index + 1:]:
            period = _period_from_text(normalize(data[date_col]))
            if period:
                # 홈택스 취소거래 금액은 원본에서 이미 음수로 제공된다.
                grouped[period] += amount(data[amount_col])
        return [ReferenceRow(y, m, source, value, filename) for (y, m), value in sorted(grouped.items())]
    raise ValueError("현금영수증 매출 제목 행을 찾지 못했습니다.")


def read_reference_file(path: str, source: str) -> list[ReferenceRow]:
    filename = Path(path).name
    if "NHN KCP" in filename.upper() and any(label in filename for label in ("계좌이체", "가상계좌")):
        source = "판매(결제)대행매출"
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return _parse_reference_pdf(path, source)
    if suffix == ".xls":
        return _parse_reference_xls(path, source)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_reference_xlsx(path, source)
    raise ValueError(f"지원하지 않는 파일 형식입니다: {suffix}")


def reference_period_label(year: int, month: int) -> str:
    return f"{year}년 {-month}분기" if month < 0 else f"{year}년 {month:02d}월"


def read_hometax_file(path: str) -> list[ReferenceRow]:
    """홈택스 카드/제로페이/판매대행 자료를 자료구분에 따라 자동 분류한다."""
    if Path(path).suffix.lower() != ".xls":
        raise ValueError("현재 제공된 홈택스 매출자료는 .xls 형식을 지원합니다.")
    ws = xlrd.open_workbook(path).sheet_by_index(0)
    values = [ws.row_values(row) for row in range(ws.nrows)]
    filename = Path(path).name

    # 월별 신용카드·제로(온누리)페이 자료
    for index, row in enumerate(values):
        headers = [normalize(value) for value in row]
        if "승인년월" in headers and "자료구분" in headers and "매출액계" in headers:
            period_col = headers.index("승인년월")
            type_col = headers.index("자료구분")
            amount_col = headers.index("매출액계")
            rows = []
            for data in values[index + 1:]:
                period = _period_from_text(normalize(data[period_col]))
                if not period:
                    continue
                data_type = normalize(data[type_col])
                source = "판매(결제)대행매출" if "판매" in data_type and "대행" in data_type else "신용카드매출"
                rows.append(ReferenceRow(period[0], period[1], source, amount(data[amount_col]), filename))
            return rows

    # 월별 판매(결제)대행업체 상세 자료 (2단 제목 행)
    for index, row in enumerate(values):
        headers = [normalize(value) for value in row]
        if "승인년월" in headers and "판매(결제)대행업체 상호" in headers:
            period_col = headers.index("승인년월")
            amount_col = next((i for i, value in enumerate(headers) if value == "계"), 3)
            grouped = defaultdict(int)
            for data in values[index + 1:]:
                period = _period_from_text(normalize(data[period_col]))
                if period:
                    grouped[period] += amount(data[amount_col])
            return [
                ReferenceRow(y, m, "판매(결제)대행매출", value, filename)
                for (y, m), value in sorted(grouped.items())
            ]

    # 분기 통합자료: 상세 월별 파일이 없을 때 조회할 수 있도록 분기 행으로 보관
    for index, row in enumerate(values):
        headers = [normalize(value) for value in row]
        if "승인분기" in headers and "자료구분" in headers and "매출액합계" in headers:
            period_col = headers.index("승인분기")
            type_col = headers.index("자료구분")
            amount_col = headers.index("매출액합계")
            rows = []
            for data in values[index + 1:]:
                match = re.search(r"(20\d{2})\D+([1-4])분기", normalize(data[period_col]))
                if not match:
                    continue
                data_type = normalize(data[type_col])
                source = "판매(결제)대행매출" if "판매" in data_type and "대행" in data_type else "신용카드매출"
                rows.append(ReferenceRow(
                    int(match.group(1)), -int(match.group(2)), source, amount(data[amount_col]), filename
                ))
            return rows
    raise ValueError("지원하는 홈택스 카드·판매대행 매출자료 형식을 찾지 못했습니다.")


def remove_redundant_quarter_rows(rows: list[ReferenceRow]) -> list[ReferenceRow]:
    """동일 분기의 월별 상세가 있으면 홈택스 분기 요약을 중복 합산하지 않는다."""
    monthly = {
        (row.year, (row.month - 1) // 3 + 1, row.source)
        for row in rows if row.month > 0
    }
    return [
        row for row in rows
        if row.month > 0 or (row.year, -row.month, row.source) not in monthly
    ]


def effective_hometax_rows(
    reference_rows: list[ReferenceRow], hometax_rows: list[ReferenceRow]
) -> list[ReferenceRow]:
    """비교용 홈택스 값: 현금은 현금매출 칸의 홈택스 현금영수증을 사용한다."""
    return [
        row for row in hometax_rows if row.source != "현금매출"
    ] + [
        row for row in reference_rows if row.source == "현금매출"
    ]


def _allocate_integer(total: int, weighted_keys: list[tuple[tuple[str, str], int]]) -> dict[tuple[str, str], int]:
    """가중치 비율로 원 단위 배분하며 결과 합계를 total과 정확히 맞춘다."""
    positive = [(key, max(0, weight)) for key, weight in weighted_keys if weight > 0]
    denominator = sum(weight for _, weight in positive)
    if not denominator:
        return {}
    sign = -1 if total < 0 else 1
    target = abs(total)
    allocated = {key: (target * weight) // denominator for key, weight in positive}
    remainder = target - sum(allocated.values())
    order = sorted(
        positive,
        key=lambda item: ((target * item[1]) % denominator, item[1], item[0]),
        reverse=True,
    )
    for index in range(remainder):
        allocated[order[index % len(order)][0]] += 1
    return {key: value * sign for key, value in allocated.items()}


def allocate_reference_sales(sls_rows: list[SalesRow], reference_rows: list[ReferenceRow]) -> list[AllocatedRow]:
    """월·매출구분별 참고자료 금액을 같은 구간의 SLS 관·항목 구성비로 배분한다."""
    weights = defaultdict(lambda: defaultdict(int))
    for row in sls_rows:
        weights[(row.year, row.month, row.source)][(row.academy, row.category)] += row.amount
    totals = defaultdict(int)
    for row in reference_rows:
        if row.month > 0:
            totals[(row.year, row.month, row.source)] += row.amount
    result = []
    missing = []
    for (year, month, source), total in sorted(totals.items()):
        source_weights = weights.get((year, month, source), {})
        if not any(value > 0 for value in source_weights.values()):
            missing.append(f"{year}년 {month:02d}월 {source}")
            continue
        allocated = _allocate_integer(total, list(source_weights.items()))
        for (academy, category), value in allocated.items():
            result.append(AllocatedRow(year, month, source, academy, category, value))
    if missing:
        raise ValueError("다음 구간은 SLS 배분기준 금액이 없습니다.\n" + "\n".join(missing))
    return result


def aggregate_allocated(
    rows: list[AllocatedRow], period: tuple[int, int] | None = None, source: str = "전체"
) -> list[dict]:
    grouped = defaultdict(lambda: defaultdict(int))
    for row in rows:
        if period and (row.year, row.month) != period:
            continue
        if source != "전체" and row.source != source:
            continue
        grouped[row.academy][row.category] += row.amount
    academies = list(ACADEMY_ORDER)
    academies += sorted(name for name in grouped if name not in ACADEMY_ORDER)
    output = []
    for academy in academies:
        record = {"학원": academy, **{category: grouped[academy][category] for category in CATEGORIES}}
        record["합계"] = sum(record[category] for category in CATEGORIES)
        output.append(record)
    return output


def build_final_tax_sales(
    sls_rows: list[SalesRow], hometax_rows: list[ReferenceRow]
) -> tuple[list[AllocatedRow], list[AllocatedRow], list[AllocatedRow]]:
    """1분기 방식으로 SLS 확정액을 빼고 남은 홈택스 매출을 40:60으로 배분한다."""
    sls_detail = [
        AllocatedRow(
            row.year, row.month, row.source, row.academy, row.category, row.amount
        )
        for row in sls_rows
    ]
    fixed_categories = {"수강료", "가맹점매출"}
    fixed_rows = [row for row in sls_detail if row.category in fixed_categories]
    fixed_totals = defaultdict(int)
    category_weights = defaultdict(lambda: defaultdict(int))
    for row in sls_rows:
        key = (row.year, row.month, row.source)
        if row.category in fixed_categories:
            fixed_totals[key] += row.amount
        elif row.category in {"가맹점유저", "교재매출"}:
            category_weights[(key, row.category)][row.academy] += max(0, row.amount)

    hometax_totals = defaultdict(int)
    for row in hometax_rows:
        if row.month > 0:
            hometax_totals[(row.year, row.month, row.source)] += row.amount

    sls_keys = {(row.year, row.month, row.source) for row in sls_rows}
    missing = [key for key in sorted(sls_keys) if key not in hometax_totals]
    if missing:
        labels = [f"{year}년 {month:02d}월 {source}" for year, month, source in missing]
        raise ValueError("다음 구간의 홈택스 매출이 없습니다.\n" + "\n".join(labels))

    remainder_rows = []
    for key in sorted(sls_keys):
        year, month, source = key
        remainder = hometax_totals[key] - fixed_totals[key]
        # 1분기 파일과 동일하게 가맹점유저 40%, 교재매출 60%로 계산한다.
        merchant_user_total = int(round(remainder * 0.4))
        category_totals = (
            ("가맹점유저", merchant_user_total),
            ("교재매출", remainder - merchant_user_total),
        )
        for category, category_total in category_totals:
            weights = [
                ((academy, category), weight)
                for academy, weight in category_weights[(key, category)].items()
            ]
            if category_total and not weights:
                raise ValueError(
                    f"{year}년 {month:02d}월 {source} {category}의 관별 SLS 배분 기준이 없습니다."
                )
            for (academy, _), value in _allocate_integer(category_total, weights).items():
                remainder_rows.append(
                    AllocatedRow(year, month, source, academy, category, value)
                )

    return sls_detail, remainder_rows, fixed_rows + remainder_rows


def aggregate(rows: list[SalesRow], period: tuple[int, int] | None, source: str) -> list[dict]:
    grouped = defaultdict(lambda: defaultdict(int))
    for row in rows:
        if period and (row.year, row.month) != period:
            continue
        if source != "전체" and row.source != source:
            continue
        grouped[row.academy][row.category] += row.amount
    result = []
    academies = list(ACADEMY_ORDER)
    academies += sorted(name for name in grouped if name not in ACADEMY_ORDER)
    for academy in academies:
        values = grouped[academy]
        record = {"학원": academy, **{key: values[key] for key in CATEGORIES}}
        record["합계"] = sum(record[key] for key in CATEGORIES)
        result.append(record)
    return result


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        PROJECTS_DIR.mkdir(parents=True, exist_ok=True)
        self.title(APP_TITLE)
        self.geometry("1080x690")
        self.minsize(900, 570)
        self.rows: list[SalesRow] = []
        self.path = ""
        self.reference_paths = {source: [] for source in REFERENCE_SOURCES}
        self.hometax_paths: list[str] = []
        self.reference_rows: list[ReferenceRow] = []
        self.hometax_rows: list[ReferenceRow] = []
        self.allocated_rows: list[AllocatedRow] = []
        self.hometax_allocated_rows: list[AllocatedRow] = []
        self.final_tax_rows: list[AllocatedRow] = []
        self._build_style()
        self._build_ui()

    def _build_style(self):
        style = ttk.Style(self)
        style.theme_use("clam")
        self.configure(bg="#17212b")
        style.configure("TFrame", background="#17212b")
        style.configure("TLabel", background="#17212b", foreground="#e6edf3", font=("맑은 고딕", 10))
        style.configure("Title.TLabel", font=("맑은 고딕", 21, "bold"), foreground="#f4f7fb")
        style.configure("Muted.TLabel", foreground="#aab8c5")
        style.configure("TButton", font=("맑은 고딕", 10), padding=(13, 8))
        style.configure("Primary.TButton", background="#377dff", foreground="white", bordercolor="#377dff")
        style.map("Primary.TButton", background=[("active", "#286be0")])
        style.configure("TCombobox", fieldbackground="#2b3948", background="#33475b", foreground="white")
        style.map("TCombobox", fieldbackground=[("readonly", "#2b3948")], foreground=[("readonly", "white")])
        style.configure("Treeview", background="#243342", fieldbackground="#243342", foreground="#e6edf3",
                        rowheight=34, font=("맑은 고딕", 10))
        style.configure("Treeview.Heading", background="#334a60", foreground="white",
                        font=("맑은 고딕", 10, "bold"), padding=8)
        style.map("Treeview", background=[("selected", "#3e6f9f")])

    def _build_ui(self):
        root = ttk.Frame(self, padding=24)
        root.pack(fill="both", expand=True)
        ttk.Label(root, text="런케이에듀 SLS 매출 조회", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            root,
            text="SLS 월별 집계표에서 런케이에듀 귀속 매출만 분류합니다. 환불이 반영된 순수납 금액 기준입니다.",
            style="Muted.TLabel",
        ).pack(anchor="w", pady=(4, 18))

        upload = ttk.Frame(root)
        upload.pack(fill="x")
        ttk.Button(upload, text="SLS 파일 첨부", command=self.choose_file, style="Primary.TButton").pack(side="left")
        ttk.Button(upload, text="부가세 참고자료 조회", command=self.open_reference_report).pack(side="left", padx=(8, 0))
        ttk.Button(
            upload, text="세무사무실 신고 최종 매출",
            command=self.open_final_tax_report, style="Primary.TButton",
        ).pack(side="left", padx=(8, 0))
        ttk.Button(upload, text="프로젝트 저장", command=self.save_project).pack(side="left", padx=(8, 0))
        ttk.Button(upload, text="프로젝트 불러오기", command=self.load_project).pack(side="left", padx=(8, 0))
        self.file_var = tk.StringVar(value="첨부된 파일이 없습니다.")
        ttk.Label(upload, textvariable=self.file_var, style="Muted.TLabel").pack(side="left", padx=12)

        filters = ttk.Frame(root)
        filters.pack(fill="x", pady=(22, 12))
        ttk.Label(filters, text="조회 월").pack(side="left")
        self.period_var = tk.StringVar(value="전체")
        self.period_box = ttk.Combobox(filters, textvariable=self.period_var, values=["전체"],
                                       state="readonly", width=15)
        self.period_box.pack(side="left", padx=(7, 22))
        ttk.Label(filters, text="매출 구분").pack(side="left")
        self.source_var = tk.StringVar(value="전체")
        self.source_box = ttk.Combobox(filters, textvariable=self.source_var,
                                       values=["전체", *SOURCES], state="readonly", width=22)
        self.source_box.pack(side="left", padx=7)
        self.period_box.bind("<<ComboboxSelected>>", self.refresh)
        self.source_box.bind("<<ComboboxSelected>>", self.refresh)

        columns = ("academy", "tuition", "merchant_user", "books", "merchant", "total")
        self.table = ttk.Treeview(root, columns=columns, show="headings")
        settings = [
            ("academy", "수납항목", 285, "w"), ("tuition", "수강료", 135, "e"),
            ("merchant_user", "가맹점유저", 135, "e"), ("books", "교재매출", 135, "e"),
            ("merchant", "가맹점매출", 135, "e"), ("total", "합계", 145, "e"),
        ]
        for key, label, width, anchor in settings:
            self.table.heading(key, text=label)
            self.table.column(key, width=width, anchor=anchor)
        self.table.pack(fill="both", expand=True)
        self.table.tag_configure("total", background="#365775", foreground="white", font=("맑은 고딕", 10, "bold"))

        footer = ttk.Frame(root)
        footer.pack(fill="x", pady=(12, 0))
        self.status_var = tk.StringVar(value="SLS 파일을 첨부해 주세요.")
        ttk.Label(footer, textvariable=self.status_var, style="Muted.TLabel").pack(side="left")
        self.export_btn = ttk.Button(footer, text="현재 조회 엑셀 저장", command=self.export, state="disabled")
        self.export_btn.pack(side="right")
        ttk.Button(footer, text="SLS 비율 배분 매출 조회", command=self.open_allocated_report).pack(
            side="right", padx=(0, 8)
        )
        ttk.Button(footer, text="홈택스 비율 배분 매출 조회", command=self.open_hometax_allocated_report).pack(
            side="right", padx=(0, 8)
        )

    def choose_file(self):
        path = filedialog.askopenfilename(title="SLS 매출 파일 선택", filetypes=FILE_TYPES)
        if not path:
            return
        try:
            rows = read_sls(path)
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"SLS 파일을 읽지 못했습니다.\n{exc}")
            return
        self.path, self.rows = path, rows
        self.allocated_rows = []
        self.hometax_allocated_rows = []
        self.final_tax_rows = []
        periods = sorted({(row.year, row.month) for row in rows})
        labels = ["전체"] + [f"{year}년 {month:02d}월" for year, month in periods]
        self.period_box.configure(values=labels)
        self.period_var.set(labels[1] if len(labels) == 2 else "전체")
        self.file_var.set(Path(path).name)
        self.export_btn.configure(state="normal")
        self.refresh()

    def save_project(self):
        if not self.path and not any(self.reference_paths.values()) and not self.hometax_paths:
            messagebox.showwarning(APP_TITLE, "저장할 첨부 파일이 없습니다.")
            return
        PROJECTS_DIR.mkdir(parents=True, exist_ok=True)
        path = filedialog.asksaveasfilename(
            title="프로젝트 저장", defaultextension=".rkproject",
            filetypes=[("런케이에듀 매출 프로젝트", "*.rkproject"), ("JSON", "*.json")],
            initialfile="런케이에듀_부가세매출.rkproject",
            initialdir=PROJECTS_DIR,
        )
        if not path:
            return
        try:
            project_path = Path(path).resolve()
            attachment_dir = project_path.parent / f"{project_path.stem}_첨부파일"

            def copy_attachment(source_path, group, index=0):
                if not source_path:
                    return ""
                source = Path(source_path).resolve()
                group_dir = attachment_dir / group
                group_dir.mkdir(parents=True, exist_ok=True)
                destination = (
                    source if source.parent == group_dir.resolve()
                    else group_dir / f"{index + 1:02d}_{source.name}"
                )
                if source != destination.resolve():
                    shutil.copy2(source, destination)
                return str(destination.relative_to(project_path.parent))

            saved_sls = copy_attachment(self.path, "SLS") if self.path else ""
            saved_reference = {
                source: [
                    copy_attachment(item, f"참고자료_{source}", index)
                    for index, item in enumerate(paths)
                ]
                for source, paths in self.reference_paths.items()
            }
            saved_hometax = [
                copy_attachment(item, "홈택스", index)
                for index, item in enumerate(self.hometax_paths)
            ]
            data = {
                "version": 2,
                "path_base": "project",
                "sls_path": saved_sls,
                "reference_paths": saved_reference,
                "hometax_paths": saved_hometax,
            }
            project_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            messagebox.showinfo(
                APP_TITLE,
                f"프로젝트와 첨부 파일 복사본을 저장했습니다.\n\n{project_path}\n{attachment_dir}",
            )
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"프로젝트를 저장하지 못했습니다.\n{exc}")

    def load_project(self):
        PROJECTS_DIR.mkdir(parents=True, exist_ok=True)
        path = filedialog.askopenfilename(
            title="프로젝트 불러오기",
            filetypes=[("런케이에듀 매출 프로젝트", "*.rkproject *.json"), ("모든 파일", "*.*")],
            initialdir=PROJECTS_DIR,
        )
        if not path:
            return
        try:
            project_path = Path(path).resolve()
            data = json.loads(project_path.read_text(encoding="utf-8"))

            def resolve_saved_path(value):
                value = normalize(value)
                if not value:
                    return ""
                candidate = Path(value)
                if data.get("path_base") == "project" and not candidate.is_absolute():
                    candidate = project_path.parent / candidate
                return str(candidate.resolve())

            sls_path = resolve_saved_path(data.get("sls_path"))
            reference_paths = data.get("reference_paths", {})
            reference_paths = {
                source: [resolve_saved_path(item) for item in reference_paths.get(source, [])]
                for source in REFERENCE_SOURCES
            }
            hometax_paths = [resolve_saved_path(item) for item in data.get("hometax_paths", [])]
            all_paths = ([sls_path] if sls_path else [])
            all_paths += [item for paths in reference_paths.values() for item in paths]
            all_paths += list(hometax_paths)
            missing = [item for item in all_paths if not Path(item).is_file()]
            if missing:
                raise ValueError("다음 첨부 파일을 찾을 수 없습니다.\n" + "\n".join(missing))

            self.path = sls_path
            self.reference_paths = {
                source: list(dict.fromkeys(reference_paths.get(source, []))) for source in REFERENCE_SOURCES
            }
            self.hometax_paths = list(dict.fromkeys(hometax_paths))
            parsed_reference = []
            for source, paths in self.reference_paths.items():
                for reference_path in paths:
                    parsed_reference.extend(read_reference_file(reference_path, source))
            parsed_hometax = []
            for hometax_path in self.hometax_paths:
                parsed_hometax.extend(read_hometax_file(hometax_path))
            self.reference_rows = parsed_reference
            self.hometax_rows = remove_redundant_quarter_rows(parsed_hometax)
            self.allocated_rows = []
            self.hometax_allocated_rows = []
            self.final_tax_rows = []
            if sls_path:
                self.rows = read_sls(sls_path)
                periods = sorted({(row.year, row.month) for row in self.rows})
                labels = ["전체"] + [f"{year}년 {month:02d}월" for year, month in periods]
                self.period_box.configure(values=labels)
                self.period_var.set("전체")
                self.file_var.set(Path(sls_path).name)
                self.export_btn.configure(state="normal")
                self.refresh()
            else:
                self.rows = []
                self.file_var.set("첨부된 SLS 파일이 없습니다.")
                self.export_btn.configure(state="disabled")
            messagebox.showinfo(
                APP_TITLE,
                "프로젝트와 첨부 파일을 불러왔습니다.\n부가세 참고자료 조회에서 바로 확인할 수 있습니다.",
            )
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"프로젝트를 불러오지 못했습니다.\n{exc}")

    def open_reference_report(self):
        window = tk.Toplevel(self)
        window.title("KSNET·KCP 부가세 참고자료 조회")
        window.geometry("1040x850")
        window.minsize(860, 680)
        frame = ttk.Frame(window, padding=20)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text="부가세 참고자료 월별 조회", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            frame,
            text="같은 세무 구분의 PDF·Excel 파일을 한 번에 여러 개 선택할 수 있습니다.",
            style="Muted.TLabel",
        ).pack(anchor="w", pady=(3, 15))

        path_vars = {source: tk.StringVar() for source in REFERENCE_SOURCES}
        hometax_var = tk.StringVar()
        upload_area = ttk.Frame(frame)
        upload_area.pack(fill="x")

        def update_path_labels():
            for source in REFERENCE_SOURCES:
                paths = self.reference_paths[source]
                if not paths:
                    path_vars[source].set("첨부 없음")
                elif len(paths) == 1:
                    path_vars[source].set(Path(paths[0]).name)
                else:
                    path_vars[source].set(f"{len(paths)}개 파일")
            if not self.hometax_paths:
                hometax_var.set("첨부 없음")
            elif len(self.hometax_paths) == 1:
                hometax_var.set(Path(self.hometax_paths[0]).name)
            else:
                hometax_var.set(f"{len(self.hometax_paths)}개 파일")
            for item in attachment_table.get_children():
                attachment_table.delete(item)
            for source in REFERENCE_SOURCES:
                for path in self.reference_paths[source]:
                    attachment_table.insert("", "end", values=(source, Path(path).name, path))
            for path in self.hometax_paths:
                attachment_table.insert("", "end", values=("홈택스 자동분류", Path(path).name, path))

        def select_files(source):
            paths = filedialog.askopenfilenames(
                title=f"{source} 참고자료 선택",
                filetypes=[("부가세 참고자료", "*.pdf *.xlsx *.xlsm *.xls"), ("모든 파일", "*.*")],
                parent=window,
            )
            if paths:
                # 같은 경로를 다시 선택해도 중복 집계하지 않는다.
                self.reference_paths[source] = list(dict.fromkeys([*self.reference_paths[source], *paths]))
                update_path_labels()
                window.lift()
                window.focus_force()

        for col, source in enumerate(REFERENCE_SOURCES):
            card = ttk.Frame(upload_area, padding=(0, 0, 12, 0))
            card.grid(row=0, column=col, sticky="nsew")
            ttk.Button(card, text=f"{source} 파일 첨부", command=lambda value=source: select_files(value)).pack(fill="x")
            ttk.Label(card, textvariable=path_vars[source], style="Muted.TLabel", wraplength=285).pack(anchor="w", pady=(5, 0))
            upload_area.columnconfigure(col, weight=1)

        def select_hometax_files():
            paths = filedialog.askopenfilenames(
                title="홈택스 매출자료 선택",
                filetypes=[("홈택스 Excel", "*.xls"), ("모든 파일", "*.*")],
                parent=window,
            )
            if paths:
                self.hometax_paths = list(dict.fromkeys([*self.hometax_paths, *paths]))
                update_path_labels()
                window.lift()
                window.focus_force()

        hometax_area = ttk.Frame(frame)
        hometax_area.pack(fill="x", pady=(12, 0))
        ttk.Button(
            hometax_area, text="홈택스 파일 첨부 · 자동분류",
            command=select_hometax_files, style="Primary.TButton",
        ).pack(side="left")
        ttk.Label(hometax_area, textvariable=hometax_var, style="Muted.TLabel", wraplength=700).pack(
            side="left", padx=10
        )

        attachment_area = ttk.Frame(frame)
        attachment_area.pack(fill="x", pady=(12, 0))
        ttk.Label(attachment_area, text="첨부된 파일 목록", font=("맑은 고딕", 11, "bold")).pack(anchor="w")
        attachment_columns = ("source", "file", "path")
        attachment_table = ttk.Treeview(
            attachment_area, columns=attachment_columns, show="headings", height=5, selectmode="extended"
        )
        for key, label, width, anchor in [
            ("source", "구분", 210, "w"),
            ("file", "파일명", 290, "w"),
            ("path", "경로", 500, "w"),
        ]:
            attachment_table.heading(key, text=label)
            attachment_table.column(key, width=width, anchor=anchor)
        attachment_table.pack(fill="x", pady=(5, 0))

        def remove_selected_files():
            selected = attachment_table.selection()
            if not selected:
                messagebox.showwarning(APP_TITLE, "삭제할 첨부 파일을 목록에서 선택해 주세요.", parent=window)
                return
            for item in selected:
                source, _, path = attachment_table.item(item, "values")
                if source == "홈택스 자동분류":
                    self.hometax_paths = [value for value in self.hometax_paths if value != path]
                else:
                    self.reference_paths[source] = [
                        value for value in self.reference_paths[source] if value != path
                    ]
            self.reference_rows = []
            self.hometax_rows = []
            self.allocated_rows = []
            self.hometax_allocated_rows = []
            self.final_tax_rows = []
            update_path_labels()
            render()

        attachment_buttons = ttk.Frame(attachment_area)
        attachment_buttons.pack(fill="x", pady=(6, 0))
        ttk.Button(
            attachment_buttons, text="선택 파일 삭제", command=remove_selected_files
        ).pack(side="right")
        update_path_labels()

        action = ttk.Frame(frame)
        action.pack(fill="x", pady=(16, 10))
        status_var = tk.StringVar(value="파일을 구분별로 첨부한 뒤 집계를 실행해 주세요.")
        ttk.Label(action, textvariable=status_var, style="Muted.TLabel").pack(side="left")

        columns = ("period", "cash", "pg", "card", "total")
        table = ttk.Treeview(frame, columns=columns, show="headings")
        settings = [
            ("period", "조회 월", 170, "center"), ("cash", "현금매출", 185, "e"),
            ("pg", "판매(결제)대행매출", 205, "e"), ("card", "신용카드매출", 185, "e"),
            ("total", "합계", 190, "e"),
        ]
        for key, label, width, anchor in settings:
            table.heading(key, text=label)
            table.column(key, width=width, anchor=anchor)
        table.pack(fill="both", expand=True)
        table.tag_configure("total", background="#365775", foreground="white", font=("맑은 고딕", 10, "bold"))

        ttk.Label(frame, text="첨부자료와 홈택스 통합 비교", font=("맑은 고딕", 11, "bold")).pack(
            anchor="w", pady=(12, 5)
        )
        compare_columns = ("source", "classified", "hometax", "difference")
        compare_table = ttk.Treeview(frame, columns=compare_columns, show="headings", height=4)
        for key, label, width in [
            ("source", "구분", 250), ("classified", "분류한 값", 210),
            ("hometax", "홈택스", 210), ("difference", "차이", 210),
        ]:
            compare_table.heading(key, text=label)
            compare_table.column(key, width=width, anchor="w" if key == "source" else "e")
        compare_table.pack(fill="x")
        compare_table.tag_configure("total", background="#365775", foreground="white",
                                    font=("맑은 고딕", 10, "bold"))

        def render():
            for item in table.get_children():
                table.delete(item)
            grouped = defaultdict(lambda: defaultdict(int))
            for row in self.reference_rows:
                grouped[(row.year, row.month)][row.source] += row.amount
            for year, month in sorted(grouped):
                values = [grouped[(year, month)][source] for source in REFERENCE_SOURCES]
                table.insert("", "end", values=(
                    reference_period_label(year, month), *(f"{value:,}" for value in values), f"{sum(values):,}"
                ))
            totals = [sum(grouped[period][source] for period in grouped) for source in REFERENCE_SOURCES]
            if grouped:
                table.insert("", "end", values=("전체 통합", *(f"{value:,}" for value in totals),
                                                f"{sum(totals):,}"), tags=("total",))
            for item in compare_table.get_children():
                compare_table.delete(item)
            classified = {source: sum(row.amount for row in self.reference_rows if row.source == source)
                          for source in REFERENCE_SOURCES}
            comparison_hometax = effective_hometax_rows(self.reference_rows, self.hometax_rows)
            hometax = {source: sum(row.amount for row in comparison_hometax if row.source == source)
                       for source in REFERENCE_SOURCES}
            for source in REFERENCE_SOURCES:
                compare_table.insert("", "end", values=(
                    source, f"{classified[source]:,}", f"{hometax[source]:,}",
                    f"{classified[source] - hometax[source]:,}",
                ))
            compare_table.insert("", "end", values=(
                "합계", f"{sum(classified.values()):,}", f"{sum(hometax.values()):,}",
                f"{sum(classified.values()) - sum(hometax.values()):,}",
            ), tags=("total",))
            file_count = sum(len(paths) for paths in self.reference_paths.values()) + len(self.hometax_paths)
            status_var.set(f"{file_count}개 파일 · {len(grouped)}개월 · 분류자료 통합 {sum(totals):,}원")

        def load_files():
            selected = [(path, source) for source, paths in self.reference_paths.items() for path in paths]
            if not selected and not self.hometax_paths:
                messagebox.showwarning(APP_TITLE, "부가세 참고자료를 하나 이상 첨부해 주세요.", parent=window)
                return
            parsed = []
            parsed_hometax = []
            errors = []
            seen = set()
            for path, source in selected:
                resolved = str(Path(path).resolve()).lower()
                if resolved in seen:
                    errors.append(f"{Path(path).name}: 다른 매출 구분에도 중복 첨부됨")
                    continue
                seen.add(resolved)
                try:
                    parsed.extend(read_reference_file(path, source))
                except Exception as exc:
                    errors.append(f"{Path(path).name}: {exc}")
            for path in self.hometax_paths:
                resolved = str(Path(path).resolve()).lower()
                if resolved in seen:
                    errors.append(f"{Path(path).name}: 다른 첨부 영역에도 중복 첨부됨")
                    continue
                seen.add(resolved)
                try:
                    parsed_hometax.extend(read_hometax_file(path))
                except Exception as exc:
                    errors.append(f"{Path(path).name}: {exc}")
            if errors:
                messagebox.showerror(
                    APP_TITLE, "일부 파일을 집계하지 못했습니다.\n\n" + "\n".join(errors), parent=window,
                )
            self.reference_rows = parsed
            self.hometax_rows = remove_redundant_quarter_rows(parsed_hometax)
            self.allocated_rows = []
            self.hometax_allocated_rows = []
            self.final_tax_rows = []
            render()

        def show_hometax_comparison():
            if not self.reference_rows:
                messagebox.showwarning(APP_TITLE, "먼저 별도 참고자료를 첨부하고 집계를 실행해 주세요.", parent=window)
                return
            if not self.hometax_rows:
                messagebox.showwarning(APP_TITLE, "홈택스 파일을 첨부하고 집계를 실행해 주세요.", parent=window)
                return

            comparison_window = tk.Toplevel(window)
            comparison_window.title("첨부자료 · 홈택스 월별 비교")
            comparison_window.geometry("980x650")
            comparison_window.minsize(780, 500)
            comparison_frame = ttk.Frame(comparison_window, padding=18)
            comparison_frame.pack(fill="both", expand=True)
            ttk.Label(comparison_frame, text="첨부자료 · 홈택스 월별 비교", style="Title.TLabel").pack(anchor="w")
            ttk.Label(
                comparison_frame,
                text="별도로 첨부한 KSNET·KCP·현금영수증 자료와 홈택스 자료를 더하지 않고 나란히 비교합니다.",
                style="Muted.TLabel",
            ).pack(anchor="w", pady=(3, 12))

            compare_columns = ("period", "source", "classified", "hometax", "difference")
            detail_table = ttk.Treeview(comparison_frame, columns=compare_columns, show="headings")
            for key, label, width, anchor in [
                ("period", "조회 월", 150, "center"), ("source", "구분", 210, "w"),
                ("classified", "별도 첨부자료", 180, "e"), ("hometax", "홈택스", 180, "e"),
                ("difference", "차이", 180, "e"),
            ]:
                detail_table.heading(key, text=label)
                detail_table.column(key, width=width, anchor=anchor)
            detail_table.pack(fill="both", expand=True)
            detail_table.tag_configure(
                "month_total", background="#314b63", foreground="white", font=("맑은 고딕", 10, "bold")
            )
            detail_table.tag_configure(
                "source_total", background="#3b6386", foreground="white", font=("맑은 고딕", 10, "bold")
            )
            detail_table.tag_configure(
                "grand_total", background="#477bab", foreground="white", font=("맑은 고딕", 10, "bold")
            )

            classified = defaultdict(int)
            hometax = defaultdict(int)
            comparison_row_keys = {}
            for row in self.reference_rows:
                classified[(row.year, row.month, row.source)] += row.amount
            comparison_hometax = effective_hometax_rows(self.reference_rows, self.hometax_rows)
            for row in comparison_hometax:
                hometax[(row.year, row.month, row.source)] += row.amount
            periods = sorted({(row.year, row.month) for row in self.reference_rows + comparison_hometax})
            grand_classified = grand_hometax = 0
            for year, month in periods:
                month_classified = month_hometax = 0
                for source in REFERENCE_SOURCES:
                    classified_value = classified[(year, month, source)]
                    hometax_value = hometax[(year, month, source)]
                    month_classified += classified_value
                    month_hometax += hometax_value
                    item_id = detail_table.insert("", "end", values=(
                        reference_period_label(year, month), source,
                        f"{classified_value:,}", f"{hometax_value:,}",
                        f"{classified_value - hometax_value:,}",
                    ))
                    comparison_row_keys[item_id] = (year, month, source)
                detail_table.insert("", "end", values=(
                    reference_period_label(year, month), "월 합계",
                    f"{month_classified:,}", f"{month_hometax:,}",
                    f"{month_classified - month_hometax:,}",
                ), tags=("month_total",))
                grand_classified += month_classified
                grand_hometax += month_hometax
            for source in REFERENCE_SOURCES:
                source_classified = sum(
                    value for (year, month, row_source), value in classified.items() if row_source == source
                )
                source_hometax = sum(
                    value for (year, month, row_source), value in hometax.items() if row_source == source
                )
                detail_table.insert("", "end", values=(
                    "전체", f"{source} 전체",
                    f"{source_classified:,}", f"{source_hometax:,}",
                    f"{source_classified - source_hometax:,}",
                ), tags=("source_total",))
            detail_table.insert("", "end", values=(
                "전체", "전체 통합",
                f"{grand_classified:,}", f"{grand_hometax:,}",
                f"{grand_classified - grand_hometax:,}",
            ), tags=("grand_total",))

            def show_difference_detail(event=None):
                if event is not None:
                    if detail_table.identify_region(event.x, event.y) != "cell":
                        return
                    if detail_table.identify_column(event.x) != "#5":
                        return
                    item_id = detail_table.identify_row(event.y)
                else:
                    selected = detail_table.selection()
                    item_id = selected[0] if selected else ""
                key = comparison_row_keys.get(item_id)
                if not key:
                    return

                year, month, source = key
                classified_rows = [
                    row for row in self.reference_rows
                    if (row.year, row.month, row.source) == key
                ]
                hometax_rows = [
                    row for row in comparison_hometax
                    if (row.year, row.month, row.source) == key
                ]
                reference_path_by_name = defaultdict(list)
                for paths in self.reference_paths.values():
                    for path in paths:
                        reference_path_by_name[Path(path).name].append(path)
                hometax_path_by_name = defaultdict(list)
                for path in self.hometax_paths:
                    hometax_path_by_name[Path(path).name].append(path)

                detail_window = tk.Toplevel(comparison_window)
                detail_window.title(f"{reference_period_label(year, month)} {source} 차이 상세")
                detail_window.geometry("980x520")
                detail_window.minsize(760, 420)
                body = ttk.Frame(detail_window, padding=18)
                body.pack(fill="both", expand=True)
                classified_total = sum(row.amount for row in classified_rows)
                hometax_total = sum(row.amount for row in hometax_rows)
                ttk.Label(
                    body,
                    text=f"{reference_period_label(year, month)} · {source}",
                    style="Title.TLabel",
                ).pack(anchor="w")
                ttk.Label(
                    body,
                    text=(
                        f"별도 첨부자료 {classified_total:,}원 - 홈택스 {hometax_total:,}원"
                        f" = {classified_total - hometax_total:,}원"
                    ),
                    style="Muted.TLabel",
                ).pack(anchor="w", pady=(3, 12))

                columns = ("side", "file", "amount", "path")
                rows_table = ttk.Treeview(body, columns=columns, show="headings")
                for column, label, width, anchor in [
                    ("side", "구분", 120, "center"),
                    ("file", "원본 파일", 260, "w"),
                    ("amount", "반영 금액", 140, "e"),
                    ("path", "파일 경로", 420, "w"),
                ]:
                    rows_table.heading(column, text=label)
                    rows_table.column(column, width=width, anchor=anchor)
                rows_table.pack(fill="both", expand=True)

                for side, rows, paths_by_name in [
                    ("별도 첨부자료", classified_rows, reference_path_by_name),
                    ("홈택스", hometax_rows, hometax_path_by_name),
                ]:
                    for row in rows:
                        paths = paths_by_name.get(row.file, [])
                        displayed_path = " / ".join(paths) if paths else row.file
                        rows_table.insert("", "end", values=(
                            side, row.file, f"{row.amount:,}", displayed_path,
                        ))
                ttk.Label(
                    body,
                    text=(
                        "원본 자료에 공통 거래 식별번호가 없어 자동 1:1 대조는 하지 않습니다. "
                        "이 창은 차이를 만든 파일별 반영 금액과 원본 경로를 보여줍니다."
                    ),
                    style="Muted.TLabel",
                ).pack(anchor="w", pady=(10, 0))
                ttk.Button(body, text="닫기", command=detail_window.destroy).pack(
                    anchor="e", pady=(10, 0)
                )

            detail_table.bind("<Double-1>", show_difference_detail)

            def export_comparison():
                path = filedialog.asksaveasfilename(
                    title="첨부자료 · 홈택스 월별 비교 저장",
                    defaultextension=".xlsx",
                    filetypes=[("Excel 통합 문서", "*.xlsx")],
                    initialfile="런케이에듀_첨부자료_홈택스_월별비교.xlsx",
                    parent=comparison_window,
                )
                if not path:
                    return
                try:
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "월별비교"
                    ws.append(["조회 월", "구분", "별도 첨부자료", "홈택스", "차이"])
                    month_total_rows = []
                    for year, month in periods:
                        month_classified = month_hometax = 0
                        for source in REFERENCE_SOURCES:
                            classified_value = classified[(year, month, source)]
                            hometax_value = hometax[(year, month, source)]
                            month_classified += classified_value
                            month_hometax += hometax_value
                            ws.append([
                                reference_period_label(year, month), source,
                                classified_value, hometax_value, classified_value - hometax_value,
                            ])
                        ws.append([
                            reference_period_label(year, month), "월 합계",
                            month_classified, month_hometax, month_classified - month_hometax,
                        ])
                        month_total_rows.append(ws.max_row)
                    source_total_rows = []
                    for source in REFERENCE_SOURCES:
                        source_classified = sum(
                            value for (*_, row_source), value in classified.items() if row_source == source
                        )
                        source_hometax = sum(
                            value for (*_, row_source), value in hometax.items() if row_source == source
                        )
                        ws.append([
                            "전체", f"{source} 전체",
                            source_classified, source_hometax, source_classified - source_hometax,
                        ])
                        source_total_rows.append(ws.max_row)
                    ws.append([
                        "전체", "전체 통합", grand_classified, grand_hometax,
                        grand_classified - grand_hometax,
                    ])
                    for cell in ws[1]:
                        cell.fill = PatternFill("solid", fgColor="1F4E78")
                        cell.font = Font(color="FFFFFF", bold=True)
                        cell.alignment = Alignment(horizontal="center")
                    for row_no in month_total_rows:
                        for cell in ws[row_no]:
                            cell.fill = PatternFill("solid", fgColor="D9EAF7")
                            cell.font = Font(bold=True)
                    for row_no in source_total_rows:
                        for cell in ws[row_no]:
                            cell.fill = PatternFill("solid", fgColor="B4C7E7")
                            cell.font = Font(bold=True)
                    for cell in ws[ws.max_row]:
                        cell.fill = PatternFill("solid", fgColor="5B9BD5")
                        cell.font = Font(color="FFFFFF", bold=True)
                    for row_no in range(2, ws.max_row + 1):
                        for col_no in range(3, 6):
                            ws.cell(row_no, col_no).number_format = '#,##0;[Red]-#,##0;\\-'
                    for index, width in enumerate((18, 28, 22, 22, 22), 1):
                        ws.column_dimensions[get_column_letter(index)].width = width
                    ws.freeze_panes = "A2"
                    wb.save(path)
                    messagebox.showinfo(APP_TITLE, "첨부자료·홈택스 월별 비교 Excel을 저장했습니다.",
                                        parent=comparison_window)
                except Exception as exc:
                    messagebox.showerror(APP_TITLE, f"비교 Excel을 저장하지 못했습니다.\n{exc}",
                                         parent=comparison_window)

            footer = ttk.Frame(comparison_frame)
            footer.pack(fill="x", pady=(10, 0))
            ttk.Label(
                footer,
                text="차이 = 별도 첨부자료 - 홈택스 · 월별 구분 행의 차이 금액을 더블클릭하면 상세를 봅니다.",
                style="Muted.TLabel",
            ).pack(side="left")
            ttk.Button(footer, text="닫기", command=comparison_window.destroy).pack(side="right")
            ttk.Button(footer, text="엑셀 저장", command=export_comparison, style="Primary.TButton").pack(
                side="right", padx=(0, 8)
            )

        def export_reference():
            if not self.reference_rows:
                messagebox.showwarning(APP_TITLE, "먼저 참고자료 집계를 실행해 주세요.", parent=window)
                return
            path = filedialog.asksaveasfilename(
                title="부가세 참고자료 집계 저장", defaultextension=".xlsx",
                filetypes=[("Excel 통합 문서", "*.xlsx")],
                initialfile="런케이에듀_부가세참고자료_월별집계.xlsx", parent=window,
            )
            if not path:
                return
            grouped = defaultdict(lambda: defaultdict(int))
            for row in self.reference_rows:
                grouped[(row.year, row.month)][row.source] += row.amount
            wb = Workbook()
            ws = wb.active
            ws.title = "월별집계"
            ws.append(["조회 월", *REFERENCE_SOURCES, "합계"])
            for year, month in sorted(grouped):
                values = [grouped[(year, month)][source] for source in REFERENCE_SOURCES]
                ws.append([reference_period_label(year, month), *values, sum(values)])
            totals = [sum(grouped[period][source] for period in grouped) for source in REFERENCE_SOURCES]
            ws.append(["전체 통합", *totals, sum(totals)])
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.font = Font(bold=True)
            for row_no in range(2, ws.max_row + 1):
                for col_no in range(2, 6):
                    ws.cell(row_no, col_no).number_format = '#,##0;[Red]-#,##0;\\-'
            for index, width in enumerate((18, 20, 24, 20, 22), 1):
                ws.column_dimensions[get_column_letter(index)].width = width

            detail = wb.create_sheet("첨부파일별")
            detail.append(["파일명", "조회 월", "매출 구분", "금액"])
            for row in sorted(self.reference_rows, key=lambda value: (value.year, value.month, value.source, value.file)):
                detail.append([row.file, reference_period_label(row.year, row.month), row.source, row.amount])
            for cell in detail[1]:
                cell.fill = PatternFill("solid", fgColor="4472C4")
                cell.font = Font(color="FFFFFF", bold=True)
            detail.column_dimensions["A"].width = 65
            detail.column_dimensions["B"].width = 18
            detail.column_dimensions["C"].width = 24
            detail.column_dimensions["D"].width = 18
            for row_no in range(2, detail.max_row + 1):
                detail.cell(row_no, 4).number_format = '#,##0;[Red]-#,##0;\\-'

            comparison = wb.create_sheet("홈택스비교")
            comparison.append(["구분", "분류한 값", "홈택스", "차이"])
            comparison_hometax = effective_hometax_rows(self.reference_rows, self.hometax_rows)
            for source in REFERENCE_SOURCES:
                classified = sum(row.amount for row in self.reference_rows if row.source == source)
                hometax = sum(row.amount for row in comparison_hometax if row.source == source)
                comparison.append([source, classified, hometax, classified - hometax])
            comparison.append([
                "합계",
                sum(row.amount for row in self.reference_rows),
                sum(row.amount for row in comparison_hometax),
                sum(row.amount for row in self.reference_rows) - sum(row.amount for row in comparison_hometax),
            ])
            for cell in comparison[1]:
                cell.fill = PatternFill("solid", fgColor="4472C4")
                cell.font = Font(color="FFFFFF", bold=True)
            for cell in comparison[comparison.max_row]:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.font = Font(bold=True)
            for row_no in range(2, comparison.max_row + 1):
                for col_no in range(2, 5):
                    comparison.cell(row_no, col_no).number_format = '#,##0;[Red]-#,##0;\\-'
            for col, width in zip("ABCD", (28, 20, 20, 20)):
                comparison.column_dimensions[col].width = width
            wb.save(path)
            messagebox.showinfo(APP_TITLE, "월별 집계와 첨부파일별 금액을 저장했습니다.", parent=window)

        ttk.Button(action, text="홈택스와 비교", command=show_hometax_comparison).pack(side="right", padx=(0, 8))
        ttk.Button(action, text="첨부자료 집계 실행", command=load_files, style="Primary.TButton").pack(side="right")
        buttons = ttk.Frame(frame)
        buttons.pack(fill="x", pady=(10, 0))
        ttk.Button(buttons, text="월별·통합 집계 엑셀 저장", command=export_reference).pack(side="left")
        ttk.Button(buttons, text="닫기", command=window.destroy).pack(side="right")
        if self.reference_rows or self.hometax_rows:
            render()

    def open_allocated_report(self):
        self._open_allocated_report("reference")

    def open_hometax_allocated_report(self):
        self._open_allocated_report("hometax")

    def open_final_tax_report(self):
        if not self.rows:
            messagebox.showwarning(APP_TITLE, "먼저 SLS 파일을 첨부해 주세요.")
            return
        effective_rows = effective_hometax_rows(self.reference_rows, self.hometax_rows)
        if not effective_rows:
            messagebox.showwarning(
                APP_TITLE,
                "먼저 부가세 참고자료 조회에서 홈택스 자료를 첨부하고 집계를 실행해 주세요.",
            )
            return
        try:
            sls_rows, remainder_rows, integrated_rows = build_final_tax_sales(
                self.rows, effective_rows
            )
            self.final_tax_rows = integrated_rows
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"최종 신고 매출을 계산하지 못했습니다.\n{exc}")
            return

        window = tk.Toplevel(self)
        window.title("세무사무실 신고 최종 매출")
        window.geometry("1120x740")
        window.minsize(900, 580)
        frame = ttk.Frame(window, padding=18)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text="세무사무실 신고 최종 매출", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            frame,
            text="SLS 수강료·가맹점매출을 먼저 반영하고, 남은 홈택스 금액을 가맹점유저 40%·교재매출 60%로 나누어 각 항목의 관별 SLS 비율로 배분합니다.",
            style="Muted.TLabel",
        ).pack(anchor="w", pady=(3, 14))

        controls = ttk.Frame(frame)
        controls.pack(fill="x", pady=(0, 10))
        ttk.Label(controls, text="조회 금액").pack(side="left")
        view_var = tk.StringVar(value="통합")
        view_box = ttk.Combobox(
            controls, textvariable=view_var, values=["SLS 금액", "남은 금액", "통합"],
            state="readonly", width=14,
        )
        view_box.pack(side="left", padx=(7, 22))
        ttk.Label(controls, text="조회 월").pack(side="left")
        periods = sorted({(row.year, row.month) for row in integrated_rows})
        period_var = tk.StringVar(value="전체")
        period_box = ttk.Combobox(
            controls, textvariable=period_var,
            values=["전체"] + [f"{year}년 {month:02d}월" for year, month in periods],
            state="readonly", width=15,
        )
        period_box.pack(side="left", padx=(7, 22))
        ttk.Label(controls, text="매출 구분").pack(side="left")
        source_var = tk.StringVar(value="전체")
        source_box = ttk.Combobox(
            controls, textvariable=source_var, values=["전체", *REFERENCE_SOURCES],
            state="readonly", width=22,
        )
        source_box.pack(side="left", padx=7)

        columns = ("academy", "tuition", "merchant_user", "books", "merchant", "total")
        table = ttk.Treeview(frame, columns=columns, show="headings")
        for key, label, width, anchor in [
            ("academy", "수납항목", 285, "w"), ("tuition", "수강료", 145, "e"),
            ("merchant_user", "가맹점유저", 145, "e"), ("books", "교재매출", 145, "e"),
            ("merchant", "가맹점매출", 145, "e"), ("total", "합계", 160, "e"),
        ]:
            table.heading(key, text=label)
            table.column(key, width=width, anchor=anchor)
        table.pack(fill="both", expand=True)
        table.tag_configure(
            "total", background="#477bab", foreground="white",
            font=("맑은 고딕", 10, "bold"),
        )
        status_var = tk.StringVar()

        def selected_period():
            match = re.match(r"(\d{4})년\s*(\d{1,2})월", period_var.get())
            return (int(match.group(1)), int(match.group(2))) if match else None

        def refresh(*_):
            rows_by_view = {
                "SLS 금액": sls_rows,
                "남은 금액": remainder_rows,
                "통합": integrated_rows,
            }
            data = aggregate_allocated(
                rows_by_view[view_var.get()], selected_period(), source_var.get()
            )
            for item in table.get_children():
                table.delete(item)
            for record in data:
                table.insert("", "end", values=(
                    record["학원"], *(f"{record[key]:,}" for key in (*CATEGORIES, "합계"))
                ))
            totals = {key: sum(record[key] for record in data) for key in (*CATEGORIES, "합계")}
            table.insert(
                "", "end",
                values=("합계금액", *(f"{totals[key]:,}" for key in (*CATEGORIES, "합계"))),
                tags=("total",),
            )
            status_var.set(
                f"{view_var.get()} · {period_var.get()} · {source_var.get()} · {totals['합계']:,}원"
            )

        for box in (view_box, period_box, source_box):
            box.bind("<<ComboboxSelected>>", refresh)
        footer = ttk.Frame(frame)
        footer.pack(fill="x", pady=(10, 0))
        ttk.Label(footer, textvariable=status_var, style="Muted.TLabel").pack(side="left")
        ttk.Button(
            footer, text="세무사 제출용 4개 엑셀 저장",
            command=lambda: self.export_allocated_workbooks(window, "final"),
            style="Primary.TButton",
        ).pack(side="right")
        ttk.Button(footer, text="닫기", command=window.destroy).pack(side="right")
        refresh()

    def _open_allocated_report(self, basis):
        if not self.rows:
            messagebox.showwarning(APP_TITLE, "먼저 메인 화면에서 SLS 파일을 첨부해 주세요.")
            return
        if not self.reference_rows:
            messagebox.showwarning(
                APP_TITLE, "먼저 부가세 참고자료 조회에서 별도 첨부자료의 집계를 실행해 주세요."
            )
            return
        if basis == "hometax" and not self.hometax_rows:
            messagebox.showwarning(
                APP_TITLE, "먼저 부가세 참고자료 조회에서 홈택스 파일을 첨부하고 집계를 실행해 주세요."
            )
            return
        allocation_source_rows = self.reference_rows
        report_name = "SLS 비율 배분 매출"
        if basis == "hometax":
            # 현금은 현금매출 칸의 홈택스 현금영수증 파일, 카드·판매대행은 홈택스 자동분류 자료.
            allocation_source_rows = [
                row for row in self.reference_rows if row.source == "현금매출"
            ] + [
                row for row in self.hometax_rows if row.source in {"신용카드매출", "판매(결제)대행매출"}
            ]
            report_name = "홈택스 비율 배분 매출"
        try:
            allocated = allocate_reference_sales(self.rows, allocation_source_rows)
            if basis == "hometax":
                self.hometax_allocated_rows = allocated
            else:
                self.allocated_rows = allocated
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"참고자료 금액을 SLS 비율로 배분하지 못했습니다.\n{exc}")
            return

        window = tk.Toplevel(self)
        window.title(f"{report_name} 조회")
        window.geometry("1120x740")
        window.minsize(900, 580)
        frame = ttk.Frame(window, padding=18)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text=f"{report_name} 조회", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            frame,
            text=(
                "홈택스 카드·판매대행과 현금매출 칸의 현금영수증 금액을 SLS 구성비로 배분했습니다."
                if basis == "hometax"
                else "월·매출구분별 참고자료 금액을 같은 월·구분의 SLS 관별·항목별 구성비로 원 단위 배분했습니다."
            ),
            style="Muted.TLabel",
        ).pack(anchor="w", pady=(3, 14))

        controls = ttk.Frame(frame)
        controls.pack(fill="x", pady=(0, 10))
        ttk.Label(controls, text="조회 월").pack(side="left")
        periods = sorted({(row.year, row.month) for row in allocated})
        period_labels = ["전체"] + [f"{year}년 {month:02d}월" for year, month in periods]
        period_var = tk.StringVar(value="전체")
        period_box = ttk.Combobox(
            controls, textvariable=period_var, values=period_labels, state="readonly", width=15
        )
        period_box.pack(side="left", padx=(7, 22))
        ttk.Label(controls, text="매출 구분").pack(side="left")
        source_var = tk.StringVar(value="전체")
        source_box = ttk.Combobox(
            controls, textvariable=source_var, values=["전체", *REFERENCE_SOURCES],
            state="readonly", width=22,
        )
        source_box.pack(side="left", padx=7)

        columns = ("academy", "tuition", "merchant_user", "books", "merchant", "total")
        table = ttk.Treeview(frame, columns=columns, show="headings")
        for key, label, width, anchor in [
            ("academy", "수납항목", 285, "w"), ("tuition", "수강료", 145, "e"),
            ("merchant_user", "가맹점유저", 145, "e"), ("books", "교재매출", 145, "e"),
            ("merchant", "가맹점매출", 145, "e"), ("total", "합계", 160, "e"),
        ]:
            table.heading(key, text=label)
            table.column(key, width=width, anchor=anchor)
        table.pack(fill="both", expand=True)
        table.tag_configure("total", background="#477bab", foreground="white",
                            font=("맑은 고딕", 10, "bold"))
        status_var = tk.StringVar()

        def selected_period():
            match = re.match(r"(\d{4})년\s*(\d{1,2})월", period_var.get())
            return (int(match.group(1)), int(match.group(2))) if match else None

        def refresh(*_):
            data = aggregate_allocated(allocated, selected_period(), source_var.get())
            for item in table.get_children():
                table.delete(item)
            for record in data:
                table.insert("", "end", values=(
                    record["학원"], *(f"{record[key]:,}" for key in (*CATEGORIES, "합계"))
                ))
            totals = {key: sum(record[key] for record in data) for key in (*CATEGORIES, "합계")}
            table.insert("", "end", values=(
                "합계금액", *(f"{totals[key]:,}" for key in (*CATEGORIES, "합계"))
            ), tags=("total",))
            exempt = totals["수강료"] + totals["교재매출"]
            taxable = totals["가맹점유저"] + totals["가맹점매출"]
            status_var.set(
                f"{period_var.get()} · {source_var.get()} · 면세 {exempt:,}원 · 과세 {taxable:,}원 · 총 {totals['합계']:,}원"
            )

        period_box.bind("<<ComboboxSelected>>", refresh)
        source_box.bind("<<ComboboxSelected>>", refresh)
        footer = ttk.Frame(frame)
        footer.pack(fill="x", pady=(10, 0))
        ttk.Label(footer, textvariable=status_var, style="Muted.TLabel").pack(side="left")
        ttk.Button(
            footer, text="세무사 제출용 4개 엑셀 저장",
            command=lambda: self.export_allocated_workbooks(window, basis),
            style="Primary.TButton",
        ).pack(side="right")
        ttk.Button(footer, text="닫기", command=window.destroy).pack(side="right", padx=(0, 8))
        refresh()

    def export_allocated_workbooks(self, parent=None, basis="reference"):
        allocated_rows = {
            "reference": self.allocated_rows,
            "hometax": self.hometax_allocated_rows,
            "final": self.final_tax_rows,
        }.get(basis, [])
        if not allocated_rows:
            messagebox.showwarning(APP_TITLE, "먼저 해당 매출 조회를 실행해 주세요.", parent=parent)
            return
        directory = filedialog.askdirectory(title="세무사 제출용 파일을 저장할 폴더 선택", parent=parent)
        if not directory:
            return

        periods = sorted({(row.year, row.month) for row in allocated_rows})
        year_text = str(periods[0][0]) if periods and len({year for year, _ in periods}) == 1 else "전체연도"
        month_text = f"{periods[0][1]}-{periods[-1][1]}월" if periods else "전체월"

        def write_table(ws, title, data, include_tax_summary=False, start_row=None):
            start = start_row or (1 if ws.max_row == 1 and ws["A1"].value is None else ws.max_row + 3)
            ws.cell(start, 1, title)
            ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
            ws.cell(start, 1).font = Font(size=15, bold=True)
            ws.cell(start, 1).alignment = Alignment(horizontal="center")
            header_row = start + 2
            for col_no, value in enumerate(["수납항목", *CATEGORIES, "합계"], 1):
                ws.cell(header_row, col_no, value)
            row_no = header_row + 1
            for record in data:
                values = [record["학원"], *(record[key] for key in (*CATEGORIES, "합계"))]
                for col_no, value in enumerate(values, 1):
                    ws.cell(row_no, col_no, value)
                row_no += 1
            totals = {key: sum(record[key] for record in data) for key in (*CATEGORIES, "합계")}
            for col_no, value in enumerate(
                ["합계금액", *(totals[key] for key in (*CATEGORIES, "합계"))], 1
            ):
                ws.cell(row_no, col_no, value)
            total_row = row_no
            for cell in ws[header_row]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            for cell in ws[total_row]:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.font = Font(bold=True)
            for row_no in range(header_row + 1, total_row + 1):
                for col_no in range(2, 7):
                    ws.cell(row_no, col_no).number_format = '#,##0;[Red]-#,##0;\\-'
            if include_tax_summary:
                exempt = totals["수강료"] + totals["교재매출"]
                taxable = totals["가맹점유저"] + totals["가맹점매출"]
                tax_start = total_row + 3
                ws.cell(tax_start, 1, "과세·면세 종합")
                ws.cell(tax_start, 1).font = Font(size=12, bold=True)
                ws.cell(tax_start + 1, 1, "구분")
                ws.cell(tax_start + 1, 2, "대상 항목")
                ws.cell(tax_start + 1, 3, "금액")
                tax_rows = [
                    ("면세", "수강료 + 교재매출", exempt),
                    ("과세", "가맹점유저 + 가맹점매출", taxable),
                    ("총 매출액", "면세 + 과세", exempt + taxable),
                ]
                for offset, values in enumerate(tax_rows, 2):
                    for col_no, value in enumerate(values, 1):
                        ws.cell(tax_start + offset, col_no, value)
                for cell in ws[tax_start + 1]:
                    if cell.column <= 3:
                        cell.fill = PatternFill("solid", fgColor="4472C4")
                        cell.font = Font(color="FFFFFF", bold=True)
                for row_no in range(tax_start + 2, tax_start + 5):
                    ws.cell(row_no, 3).number_format = '#,##0;[Red]-#,##0;\\-'
                ws.cell(tax_start + 4, 1).font = Font(bold=True)
                ws.cell(tax_start + 4, 3).font = Font(bold=True)
            if start == 1:
                ws.freeze_panes = "A4"
            for index, width in enumerate((31, 19, 19, 19, 19, 21), 1):
                ws.column_dimensions[get_column_letter(index)].width = width

        def create_workbook(source):
            wb = Workbook()
            wb.remove(wb.active)
            label = "합산액 매출" if source == "전체" else source
            summary = aggregate_allocated(allocated_rows, None, source)
            ws = wb.create_sheet("종합")
            write_table(ws, f"{year_text}년 {month_text} 런케이에듀 {label} 종합", summary, source == "전체")
            for year, month in periods:
                monthly = aggregate_allocated(allocated_rows, (year, month), source)
                write_table(
                    ws, f"{year}년 {month:02d}월 런케이에듀 {label}",
                    monthly, False,
                )
            for year, month in periods:
                monthly = aggregate_allocated(allocated_rows, (year, month), source)
                ws = wb.create_sheet(f"{month}월")
                write_table(
                    ws, f"{year}년 {month:02d}월 런케이에듀 {label}",
                    monthly, source == "전체",
                )
            return wb

        created = []
        try:
            targets = [*REFERENCE_SOURCES, "전체"]
            basis_label = {
                "hometax": "홈택스기준",
                "reference": "첨부자료기준",
                "final": "세무사무실신고최종",
            }[basis]
            for source in targets:
                label = "합산액매출" if source == "전체" else source
                filename = f"런케이에듀_{year_text}년_{month_text}_{basis_label}_{label}_세무사제출용.xlsx"
                path = Path(directory) / filename
                create_workbook(source).save(path)
                created.append(path.name)
            messagebox.showinfo(
                APP_TITLE,
                "세무사 제출용 Excel 4개를 저장했습니다.\n\n" + "\n".join(created),
                parent=parent,
            )
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"배분 매출 Excel을 저장하지 못했습니다.\n{exc}", parent=parent)

    def selected_period(self) -> tuple[int, int] | None:
        match = re.match(r"(\d{4})년\s*(\d{1,2})월", self.period_var.get())
        return (int(match.group(1)), int(match.group(2))) if match else None

    def current_data(self) -> list[dict]:
        return aggregate(self.rows, self.selected_period(), self.source_var.get())

    def refresh(self, *_):
        for item in self.table.get_children():
            self.table.delete(item)
        if not self.rows:
            return
        data = self.current_data()
        for record in data:
            self.table.insert("", "end", values=[
                record["학원"], *(f"{record[key]:,}" for key in (*CATEGORIES, "합계"))
            ])
        totals = {key: sum(record[key] for record in data) for key in (*CATEGORIES, "합계")}
        self.table.insert("", "end", values=["합계금액", *(f"{totals[key]:,}" for key in (*CATEGORIES, "합계"))],
                          tags=("total",))
        self.status_var.set(
            f"{self.period_var.get()} · {self.source_var.get()} · 런케이에듀 매출 {totals['합계']:,}원"
        )

    def export(self):
        data = self.current_data()
        period, source = self.period_var.get(), self.source_var.get()
        initial = re.sub(r'[\\/:*?"<>|]', "_", f"{period}_{source}_런케이에듀_매출.xlsx")
        path = filedialog.asksaveasfilename(
            title="매출 조회 결과 저장", defaultextension=".xlsx",
            filetypes=[("Excel 통합 문서", "*.xlsx")], initialfile=initial,
        )
        if not path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "매출조회"
            ws.append([f"{period} 런케이에듀 {source}"])
            ws.merge_cells("A1:F1")
            ws["A1"].font = Font(size=15, bold=True)
            ws["A1"].alignment = Alignment(horizontal="center")
            ws.append([])
            headers = ["수납항목", *CATEGORIES, "합계"]
            ws.append(headers)
            for record in data:
                ws.append([record["학원"], *(record[key] for key in (*CATEGORIES, "합계"))])
            ws.append(["합계금액", *(sum(record[key] for record in data) for key in (*CATEGORIES, "합계"))])
            for cell in ws[3]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.font = Font(bold=True)
            for row_no in range(4, ws.max_row + 1):
                for col_no in range(2, 7):
                    ws.cell(row_no, col_no).number_format = '#,##0;[Red]-#,##0;\\-'
            for index, width in enumerate((30, 18, 18, 18, 18, 20), 1):
                ws.column_dimensions[get_column_letter(index)].width = width
            ws.freeze_panes = "A4"
            wb.save(path)
            messagebox.showinfo(APP_TITLE, "매출 조회 결과를 저장했습니다.")
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"엑셀 파일을 저장하지 못했습니다.\n{exc}")


if __name__ == "__main__":
    App().mainloop()
